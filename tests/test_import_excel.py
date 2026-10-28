"""Importazione di clienti e polizze da Excel e CSV."""
import io
import os
import tempfile
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from gestionale import create_app
from gestionale import import_excel as ie
from gestionale.models import Client, Policy, db


@pytest.fixture
def app():
    db_fd, db_path = tempfile.mkstemp(suffix=".db")
    app = create_app({
        "TESTING": True,
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{db_path}",
        "UPLOAD_FOLDER": tempfile.mkdtemp(),
        "SECRET_KEY": "test",
    })
    yield app
    os.close(db_fd)
    os.unlink(db_path)


@pytest.fixture
def client(app):
    return app.test_client()


def _xlsx(fogli: dict) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    for nome, righe in fogli.items():
        foglio = libro.create_sheet(nome)
        for riga in righe:
            foglio.append(riga)
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()


def test_modello_ha_i_fogli_e_si_reimporta(app):
    dati = ie.modello_xlsx()
    libro = load_workbook(io.BytesIO(dati))
    assert {"Clienti", "Polizze", "Valori", "Istruzioni"} <= set(libro.sheetnames)
    with app.app_context():
        a = ie.analizza(dati, "modello.xlsx")
        # La riga di esempio: un cliente nuovo e una polizza collegata a lui.
        assert a["conteggi"]["clienti_nuovi"] == 1
        assert a["conteggi"]["polizze_nuove"] == 1
        assert a["polizze"][0]["dati"]["ramo"] == "rct_rco"
        assert a["polizze"][0]["dati"]["tacito_rinnovo"] is True
        assert a["polizze"][0]["nome_cliente"] == "Rossi Impianti S.r.l."


def test_intestazioni_libere_e_conversioni(app):
    dati = _xlsx({"Foglio1": [
        ["Elenco clienti 2026"],  # titolo sopra l'intestazione
        ["Nominativo", "P.IVA", "Comune", "CAP", "Tel", "Fatturato"],
        ["Bianchi Srl", 1234567890, "Milano", 2100, "02 123", "1.250.000,50"],
        ["", None, None, None, None, None],
        ["Verdi Mario", None, "Roma", "00184", None, None],
    ]})
    with app.app_context():
        a = ie.analizza(dati, "miei clienti.xlsx")
        assert a["conteggi"]["clienti_nuovi"] == 2
        b = a["clienti"][0]["dati"]
        assert b["partita_iva"] == "01234567890"  # zero iniziale ripristinato
        assert b["cap"] == "02100"
        assert b["fatturato_annuo"] == 1250000.50
        assert a["clienti"][0]["riga"] == 3


def test_polizze_creano_clienti_mancanti_e_saltano_doppioni(app):
    with app.app_context():
        c = Client(tipo="azienda", ragione_sociale="Esistente SpA", partita_iva="11111111111")
        db.session.add(c)
        db.session.flush()
        db.session.add(Policy(client_id=c.id, compagnia="Generali", numero_polizza="A-1"))
        db.session.commit()
        cid = c.id

    dati = _xlsx({"Polizze": [
        ["Cliente", "Compagnia", "N. polizza", "Ramo", "Premio", "Scadenza", "Provvigione"],
        ["11111111111", "Generali", "A-1", "RCT", 100, date(2026, 12, 31), 0.15],   # già presente
        ["esistente spa", "Allianz", "B-2", "Infortuni", "1.200,00", "31/12/2026", "12%"],
        ["Nuovo Cliente", "Unipol", "C-3", "ramo strano", None, "non è una data", None],
        ["Nuovo Cliente", "Unipol", "C-4", "TCM", None, None, None],
        ["", "Unipol", "", None, None, None, None],
    ]})
    with app.app_context():
        a = ie.analizza(dati, "portafoglio.xlsx")
        esiti = [p["esito"] for p in a["polizze"]]
        assert esiti == ["esistente", "nuova", "errore", "nuova", "errore"]
        assert a["polizze"][1]["dati"]["provvigione_perc"] == 12
        assert a["polizze"][1]["nome_cliente"] == "Esistente SpA"
        assert any("non riconosciuto" in w for w in a["polizze"][2]["avvisi"])
        assert a["polizze"][3]["dati"]["ramo"] == "vita"
        assert a["conteggi"]["clienti_nuovi"] == 1  # «Nuovo Cliente», una volta sola

        esito = ie.importa(a)
        assert esito == {"clienti": 1, "completati": 0, "polizze": 2, "persone": 0}
        assert Policy.query.filter_by(client_id=cid).count() == 2
        nuovo = Client.query.filter_by(ragione_sociale="Nuovo Cliente").one()
        assert [p.numero_polizza for p in nuovo.polizze] == ["C-4"]


def test_clienti_esistenti_completati_non_sovrascritti(app):
    with app.app_context():
        db.session.add(Client(tipo="azienda", ragione_sociale="Rossi", partita_iva="22222222222", citta="Brescia"))
        db.session.commit()
        dati = _xlsx({"Clienti": [
            ["Ragione sociale", "Partita IVA", "Città", "Email"],
            ["Rossi Srl (nome diverso)", "22222222222", "Milano", "info@rossi.it"],
        ]})
        a = ie.analizza(dati, "x.xlsx")
        assert a["clienti"][0]["esito"] == "esistente"
        assert ie.importa(a)["completati"] == 1
        c = Client.query.one()
        assert (c.ragione_sociale, c.citta, c.email) == ("Rossi", "Brescia", "info@rossi.it")


def test_csv_con_punto_e_virgola(app):
    testo = "Ragione sociale;Codice fiscale;Provincia\nNeri Luca;nrilcu80a01h501x;rm\n".encode("cp1252")
    with app.app_context():
        a = ie.analizza(testo, "export.csv")
        d = a["clienti"][0]["dati"]
        assert d["codice_fiscale"] == "NRILCU80A01H501X" and d["provincia"] == "RM"


def test_file_non_validi(app):
    with app.app_context():
        with pytest.raises(ie.FileNonValido, match="xlsx"):
            ie.analizza(b"vecchio", "vecchio.xls")
        with pytest.raises(ie.FileNonValido):
            ie.analizza(b"non sono uno zip", "rotto.xlsx")
        with pytest.raises(ie.FileNonValido, match="colonne"):
            ie.analizza(_xlsx({"Foglio": [["Colore", "Peso"], ["rosso", 3]]}), "a.xlsx")


def test_flusso_web_anteprima_e_conferma(app, client):
    assert client.get("/clienti/importa/").status_code == 200
    r = client.get("/clienti/importa/modello.xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"

    r = client.post("/clienti/importa/", data={"file": (io.BytesIO(r.data), "modello.xlsx")},
                    content_type="multipart/form-data")
    assert r.status_code == 200
    assert "Anteprima" in r.get_data(as_text=True)
    with app.app_context():
        assert Client.query.count() == 0  # l'anteprima non scrive nulla

    with client.session_transaction() as s:
        token = s["dc_import"]["token"]
    assert client.post("/clienti/importa/conferma", data={"token": "sbagliato"}).status_code == 302
    r = client.post("/clienti/importa/conferma", data={"token": token}, follow_redirects=True)
    assert "Importazione completata" in r.get_data(as_text=True)
    with app.app_context():
        assert Client.query.count() == 1 and Policy.query.count() == 1
    # Il file temporaneo è stato cancellato: una seconda conferma non reimporta.
    client.post("/clienti/importa/conferma", data={"token": token})
    with app.app_context():
        assert Policy.query.count() == 1


def test_cliente_esistente_con_nome_diverso_non_si_duplica(app):
    with app.app_context():
        db.session.add(Client(tipo="azienda", ragione_sociale="ROSSI IMPIANTI", partita_iva="01234567890"))
        db.session.commit()
        a = ie.analizza(ie.modello_xlsx(), "modello.xlsx")
        assert a["conteggi"]["clienti_nuovi"] == 0 and a["conteggi"]["clienti_esistenti"] == 1
        assert a["polizze"][0]["nome_cliente"] == "ROSSI IMPIANTI"
        assert a["ignorati"] == []  # i fogli di aiuto del modello non sono segnalati
        ie.importa(a)
        assert Client.query.count() == 1 and Policy.query.count() == 1


def test_stesso_nome_codici_diversi_non_si_uniscono(app):
    with app.app_context():
        db.session.add(Client(tipo="dipendente", ragione_sociale="Mario Rossi", codice_fiscale="RSSMRA80A01H501U"))
        db.session.commit()
        dati = _xlsx({
            "Clienti": [["Nome", "Codice fiscale"], ["Mario Rossi", "RSSMRA75B02F205X"]],
            "Polizze": [["Cliente", "Codice fiscale", "Compagnia", "Numero polizza"],
                        ["Mario Rossi", "RSSMRA75B02F205X", "Generali", "Z-1"]],
        })
        a = ie.analizza(dati, "omonimi.xlsx")
        assert a["clienti"][0]["esito"] == "nuovo"
        assert any("stesso nome" in w for w in a["clienti"][0]["avvisi"])
        ie.importa(a)
        vecchio = Client.query.filter_by(codice_fiscale="RSSMRA80A01H501U").one()
        nuovo = Client.query.filter_by(codice_fiscale="RSSMRA75B02F205X").one()
        assert vecchio.polizze == [] and [p.numero_polizza for p in nuovo.polizze] == ["Z-1"]


def test_righe_di_esempio_segnalate(app):
    with app.app_context():
        a = ie.analizza(ie.modello_xlsx(), "modello.xlsx")
        assert any("esempio" in w for w in a["clienti"][0]["avvisi"])
        assert any("esempio" in w for w in a["polizze"][0]["avvisi"])


def test_persone_con_nome_cognome_albo_e_residenza(app):
    from gestionale.models import PersonaCollegata  # noqa: F401

    dati = _xlsx({"Clienti": [
        ["Nome", "Cognome", "Codice fiscale", "Partita IVA", "Albo", "N. iscrizione", "Nato il", "Comune di residenza"],
        ["Giulia", "Bianchi", "BNCGLI80D52B157X", "01234567891", "Ordine degli Ingegneri", "A123", "12/04/1980", "Monza"],
        ["Anna", "Neri", "NRENNA85A41F205Y", None, None, None, None, None],
    ]})
    with app.app_context():
        a = ie.analizza(dati, "persone.xlsx")
        g, n = a["clienti"][0]["dati"], a["clienti"][1]["dati"]
        assert g["ragione_sociale"] == "Giulia Bianchi" and g["tipo"] == "professionista"
        assert n["tipo"] == "dipendente" and n["ragione_sociale"] == "Anna Neri"
        ie.importa(a)
        c = Client.query.filter_by(cognome="Bianchi").one()
        assert (c.albo, c.albo_numero, c.data_nascita, c.residenza_citta) == (
            "Ordine degli Ingegneri", "A123", date(1980, 4, 12), "Monza")


def test_foglio_amministratori(app):
    from gestionale.models import PersonaCollegata

    with app.app_context():
        c = Client(tipo="azienda", ragione_sociale="Alfa S.r.l.", partita_iva="11111111111")
        db.session.add(c)
        db.session.flush()
        db.session.add(PersonaCollegata(client_id=c.id, ruolo="consigliere", nome="Già", cognome="Presente"))
        db.session.commit()
    dati = _xlsx({"Amministratori": [
        ["Società", "Carica", "Nome", "Cognome", "CF", "Quota"],
        ["11111111111", "A.D.", "Luca", "Verdi", "vrdlcu70a01b157x", None],
        ["Alfa S.r.l.", "Socio", "Sara", "Gialli", None, 0.3],
        ["Alfa S.r.l.", "Direttore tecnico", "Piero", "Blu", None, None],
        ["Alfa S.r.l.", "Consigliere", "Già", "Presente", None, None],
        ["Beta S.p.A.", "Presidente", "Ugo", "Rosa", None, None],
        ["Alfa S.r.l.", "Socio", "Senza", "", None, None],
    ]})
    with app.app_context():
        a = ie.analizza(dati, "cda.xlsx")
        esiti = [r["esito"] for r in a["persone"]]
        assert esiti == ["nuova", "nuova", "nuova", "esistente", "nuova", "errore"]
        ruoli = [(r["dati"]["ruolo"], r["dati"]["ruolo_altro"]) for r in a["persone"][:3]]
        assert ruoli == [("amministratore_delegato", None), ("socio", None), ("altro", "Direttore tecnico")]
        assert a["persone"][1]["dati"]["quota_perc"] == 30
        assert a["conteggi"]["clienti_nuovi"] == 1  # «Beta S.p.A.» creato dal foglio amministratori
        esito = ie.importa(a)
        assert esito["persone"] == 4
        beta = Client.query.filter_by(ragione_sociale="Beta S.p.A.").one()
        assert [p.ruolo_label for p in beta.persone] == ["Presidente del CdA"]
        alfa = Client.query.filter_by(partita_iva="11111111111").one()
        assert {p.ruolo_label for p in alfa.persone} == {"Consigliere", "Amministratore delegato", "Socio",
                                                       "Direttore tecnico"}
        assert next(p for p in alfa.persone if p.cognome == "Verdi").codice_fiscale == "VRDLCU70A01B157X"


def test_modello_contiene_il_foglio_amministratori(app):
    libro = load_workbook(io.BytesIO(ie.modello_xlsx()))
    assert "Amministratori" in libro.sheetnames
    with app.app_context():
        a = ie.analizza(ie.modello_xlsx(), "modello.xlsx")
        assert a["conteggi"]["persone_nuove"] == 1
        assert a["persone"][0]["nome_cliente"] == ie.ESEMPIO_CLIENTE
        assert any("esempio" in w for w in a["persone"][0]["avvisi"])
